#!/usr/bin/env python3
"""Force every paragraph inside a DOCX table cell to use zero first-line indent."""

from __future__ import annotations

import argparse
import os
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def _iter_table_paragraphs(container):
    """Yield paragraphs in top-level and nested tables for a document story."""
    seen_cells = set()
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_id = id(cell._tc)
                if cell_id in seen_cells:
                    continue
                seen_cells.add(cell_id)
                yield from cell.paragraphs
                yield from _iter_table_paragraphs(cell)


def iter_table_paragraphs(doc):
    """Yield table-cell paragraphs from the body and all distinct headers/footers."""
    seen_parts = set()
    roots = [doc]
    for section in doc.sections:
        roots.extend(
            [
                section.header,
                section.first_page_header,
                section.even_page_header,
                section.footer,
                section.first_page_footer,
                section.even_page_footer,
            ]
        )

    for root in roots:
        part = getattr(root, "part", None)
        part_key = getattr(part, "partname", None) or id(root)
        if part_key in seen_parts:
            continue
        seen_parts.add(part_key)
        yield from _iter_table_paragraphs(root)


def set_zero_first_line_indent(paragraph):
    """Apply an explicit zero first-line indent without changing other formatting."""
    paragraph.paragraph_format.first_line_indent = Pt(0)
    p_pr = paragraph._p.get_or_add_pPr()
    ind = p_pr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        p_pr.append(ind)
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:firstLineChars"), "0")
    ind.attrib.pop(qn("w:hanging"), None)
    ind.attrib.pop(qn("w:hangingChars"), None)


def normalize_table_paragraphs(doc):
    """Normalize all table paragraphs and return the number processed."""
    count = 0
    for paragraph in iter_table_paragraphs(doc):
        set_zero_first_line_indent(paragraph)
        count += 1
    return count


def table_indent_issues(doc):
    """Return descriptions of table paragraphs lacking an explicit zero indent."""
    issues = []
    for index, paragraph in enumerate(iter_table_paragraphs(doc), start=1):
        p_pr = paragraph._p.pPr
        ind = p_pr.find(qn("w:ind")) if p_pr is not None else None
        valid = (
            ind is not None
            and ind.get(qn("w:firstLine")) == "0"
            and ind.get(qn("w:firstLineChars")) == "0"
            and ind.get(qn("w:hanging")) is None
            and ind.get(qn("w:hangingChars")) is None
        )
        if not valid:
            preview = paragraph.text.replace("\n", " ")[:40]
            issues.append(f"table paragraph {index}: {preview!r}")
    return issues


def save_document(doc, destination):
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=f".{destination.stem}.", suffix=".docx", dir=destination.parent, delete=False
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        doc.save(temporary_path)
        os.replace(temporary_path, destination)
    finally:
        temporary_path.unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(
        description="Remove first-line indentation from every paragraph inside DOCX tables."
    )
    parser.add_argument("input", type=Path, help="Input DOCX file")
    parser.add_argument("--output", type=Path, help="Output DOCX; defaults to in-place update")
    parser.add_argument(
        "--check",
        action="store_true",
        help="Check only; exit non-zero if any table paragraph is not explicitly zero-indented",
    )
    args = parser.parse_args()

    doc = Document(args.input)
    if args.check:
        issues = table_indent_issues(doc)
        if issues:
            for issue in issues[:20]:
                print(issue)
            if len(issues) > 20:
                print(f"... and {len(issues) - 20} more")
            raise SystemExit(1)
        print(f"OK: all table paragraphs have explicit zero first-line indent: {args.input}")
        return

    count = normalize_table_paragraphs(doc)
    destination = args.output or args.input
    save_document(doc, destination)
    print(f"Normalized {count} table paragraphs: {destination}")


if __name__ == "__main__":
    main()
