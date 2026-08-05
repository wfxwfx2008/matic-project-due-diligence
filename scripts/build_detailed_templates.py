#!/usr/bin/env python3
"""Build MATIC detailed due-diligence DOCX templates."""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from build_template import (
    CN_BODY,
    CN_HEI,
    CN_KAI,
    CN_TITLE,
    add_page_number,
    add_placeholder_table,
    configure_style,
    enable_field_updates,
    get_or_add_paragraph_style,
    set_cell_text,
    set_run_font,
)
from normalize_docx_tables import normalize_table_paragraphs, set_zero_first_line_indent


ROOT = Path(__file__).resolve().parents[1]
STARTUP_OUTPUT = ROOT / "assets" / "matic-detailed-dd-startup-template.docx"
REPORT_OUTPUT = ROOT / "assets" / "matic-detailed-due-diligence-template.docx"


def new_matic_document(justify_body=False):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.2)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)

    body_alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify_body else None
    configure_style(doc.styles["Normal"], CN_BODY, 16, align=body_alignment, indent=True)
    styles = {
        "title": get_or_add_paragraph_style(doc, "MATIC Title"),
        "subtitle": get_or_add_paragraph_style(doc, "MATIC Subtitle"),
        "heading1": get_or_add_paragraph_style(doc, "MATIC Heading 1"),
        "heading2": get_or_add_paragraph_style(doc, "MATIC Heading 2"),
        "heading3": get_or_add_paragraph_style(doc, "MATIC Heading 3"),
        "caption": get_or_add_paragraph_style(doc, "MATIC Caption"),
    }
    configure_style(styles["title"], CN_TITLE, 22, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    configure_style(styles["subtitle"], CN_KAI, 16, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    configure_style(styles["heading1"], CN_HEI, 16, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, outline_level=0)
    configure_style(styles["heading2"], CN_KAI, 16, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, outline_level=1)
    configure_style(styles["heading3"], CN_KAI, 16, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, outline_level=2)
    configure_style(styles["caption"], CN_BODY, 10.5, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)

    source = (
        doc.styles["Source Text"]
        if "Source Text" in [style.name for style in doc.styles]
        else doc.styles.add_style("Source Text", WD_STYLE_TYPE.PARAGRAPH)
    )
    configure_style(source, CN_BODY, 10.5, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False)

    for style in (styles["heading1"], styles["heading2"], styles["heading3"]):
        style.paragraph_format.keep_with_next = True

    add_page_number(section.footer.paragraphs[0])
    enable_field_updates(doc)
    return doc


def add_meta_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for table_row, (label, value) in zip(table.rows, rows):
        set_cell_text(table_row.cells[0], label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table_row.cells[1], value)
    return table


def add_cover(doc, document_name, meta_rows):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_zero_first_line_indent(p)
    p.paragraph_format.space_after = Pt(48)
    run = p.add_run("内部研判材料，仅供决策参考")
    set_run_font(run, CN_KAI, 16)

    p = doc.add_paragraph(style="MATIC Title")
    set_zero_first_line_indent(p)
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(24)
    p.add_run("[项目名称]")
    p = doc.add_paragraph(style="MATIC Subtitle")
    set_zero_first_line_indent(p)
    p.add_run(document_name)
    doc.add_paragraph("")
    add_meta_table(doc, meta_rows)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(84)
    set_zero_first_line_indent(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("长三角医学先进技术创新中心")
    set_run_font(run, CN_BODY, 16)
    p = doc.add_paragraph()
    set_zero_first_line_indent(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("[YYYY年MM月DD日]")
    set_run_font(run, CN_BODY, 16)
    run.add_break(WD_BREAK.PAGE)


def add_evidence_item(doc, number, text, evidence):
    doc.add_paragraph(f"{number}. {text}")
    p = doc.add_paragraph(style="Source Text")
    p.add_run(evidence)


def add_strength_risk_page(doc, heading, strength_count=4, risk_count=3):
    doc.add_paragraph(heading, style="MATIC Heading 2")
    doc.add_paragraph("（一）主要优势", style="MATIC Heading 3")
    for idx in range(1, strength_count + 1):
        add_evidence_item(
            doc,
            idx,
            "[填写经详细尽调确认、对决策有实质影响且不与其他板块重复的优势。]",
            "（支持：[专家姓名]〔类别〕；访谈纪要[INT-XX]）",
        )
    doc.add_paragraph("（二）主要风险", style="MATIC Heading 3")
    for idx in range(1, risk_count + 1):
        add_evidence_item(
            doc,
            idx,
            "[填写经详细尽调识别、对推进条件有实质影响且不与其他板块重复的风险。]",
            "（提出风险：[专家姓名]〔类别〕；访谈纪要[INT-XX]；[如有不同意见及补证要求]）",
        )


def add_second_part_section(doc, heading, kind="text"):
    doc.add_paragraph(heading, style="MATIC Heading 2")
    if kind == "scope":
        add_placeholder_table(
            doc,
            ["证据类别", "数量", "覆盖问题", "主要限制", "检索或访谈截止日期"],
            4,
        )
    elif kind == "issues":
        add_placeholder_table(
            doc,
            ["问题编号", "核心问题", "标准尽调判断", "新增证据", "当前判断", "状态"],
            4,
        )
    elif kind == "conclusions":
        add_placeholder_table(
            doc,
            ["结论编号", "优势/风险", "当前结论", "专家证据", "客观证据", "状态"],
            4,
        )
    elif kind == "disputes":
        add_placeholder_table(
            doc,
            ["分歧编号", "主题", "支持观点", "反对观点", "适用条件", "补充调研", "状态"],
            3,
        )
    elif kind == "risks":
        add_placeholder_table(
            doc,
            ["风险", "类别", "证据", "严重性", "可能性", "状态", "解除/验证条件"],
            4,
        )
    elif kind == "plan":
        add_placeholder_table(
            doc,
            ["优先级", "待补材料或问题", "责任对象", "验证方式", "预期形成的判断"],
            4,
        )
    elif kind == "sources":
        add_placeholder_table(
            doc,
            ["序号", "机构或作者", "标题/材料名称", "日期", "访问日期", "直达链接或材料编号"],
            5,
        )
        doc.add_paragraph("附录：专家访谈及材料文件索引", style="MATIC Heading 3")
        add_placeholder_table(
            doc,
            ["材料编号", "访谈编号", "专家/提供方", "文件名称", "日期", "使用范围"],
            4,
        )
    elif kind == "competition":
        doc.add_paragraph("（一）竞争格局、代表产品与替代路线", style="MATIC Heading 3")
        add_placeholder_table(
            doc,
            ["对比编号/对象", "类别与阶段", "对比维度", "本项目相对优劣势", "证据/可比性", "竞争影响"],
            6,
        )
        doc.add_paragraph("（二）专家证据、持续性与验证计划", style="MATIC Heading 3")
        add_placeholder_table(
            doc,
            ["当前判断", "专家支持/反对及限制", "持续性或可弥补性", "验证动作与判定标准"],
            4,
        )
        doc.add_paragraph(
            "[第二部分完整保留全部有效比较；仅有公开证据或AI推断的内容标注待专家验证，不得冒充第一部分专家结论。]"
        )
    elif kind == "market":
        doc.add_paragraph("（一）五层市场模型及变化", style="MATIC Heading 3")
        add_placeholder_table(
            doc,
            ["市场层级", "边界与公式", "当前区间", "相较标准尽调的变化", "依据/可靠性/敏感变量"],
            5,
        )
        doc.add_paragraph("（二）商业验证与项目销售校验", style="MATIC Heading 3")
        doc.add_paragraph(
            "[说明中国市场为主、全球市场为补充；校验SOM、项目销售、产能、渠道和医院准入。无直接数据时披露代理变量与调整依据。]"
        )
        doc.add_paragraph("（三）完整市场结论及状态", style="MATIC Heading 3")
        add_placeholder_table(
            doc,
            ["结论编号", "优势/风险", "当前结论", "专家证据", "客观证据", "状态"],
            4,
        )
    else:
        doc.add_paragraph(
            "[围绕关键问题展开完整分析，区分项目方陈述、专家判断、公开事实、第三方证据、独立推断和待核实事项。]"
        )
        doc.add_paragraph("（一）关键问题与证据", style="MATIC Heading 3")
        doc.add_paragraph("[填写问题、证据、分歧、适用条件和证据局限。]")
        doc.add_paragraph("（二）完整结论及状态", style="MATIC Heading 3")
        doc.add_paragraph("[完整保留本专项有效结论，并与工作台编号一致。]")


def build_startup_template():
    doc = new_matic_document()
    add_cover(
        doc,
        "详细尽调启动方案",
        [
            ("项目名称", "[项目名称]"),
            ("方案版本", "V0.1"),
            ("当前阶段", "启动详细尽调"),
            ("标准报告版本", "[版本及日期]"),
            ("方案日期", "[YYYY年MM月DD日]"),
        ],
    )

    doc.add_paragraph("一、项目及继承判断", style="MATIC Heading 1")
    doc.add_paragraph("[简述项目，并列明从标准尽调中继承的战略性、先进性、核心优势、主要风险和推进前置条件。]")
    add_placeholder_table(doc, ["事项", "标准尽调判断", "当前证据", "详细尽调处理方式"], 5)

    doc.add_paragraph("二、详细尽调关键问题", style="MATIC Heading 1")
    add_placeholder_table(
        doc,
        ["问题编号", "核心问题", "影响的决策", "现有证据", "证据缺口", "优先级", "验证方式"],
        6,
    )

    doc.add_paragraph("三、专项调研与补证方案", style="MATIC Heading 1")
    add_placeholder_table(
        doc,
        ["专项", "尽调目标", "核心问题", "拟采用证据", "阶段交付", "完成条件"],
        6,
    )
    doc.add_paragraph(
        "说明：继承标准尽调的竞品与替代方案比较，建立暂定竞品矩阵并标注待专家验证；继承五层市场模型，市场规模不作为固定专家访谈重点。"
    )

    doc.add_paragraph("四、调研对象与访谈安排", style="MATIC Heading 1")
    doc.add_paragraph(
        "[候选库原则上建立30—40名或家以上对象，优先江浙沪和长三角地区；实际访谈约15—20人，各类别不平均分配。]"
    )
    add_placeholder_table(
        doc,
        ["类别", "建议人数", "第一轮对象", "拟验证问题", "独立性/利益冲突", "优先级"],
        6,
    )

    doc.add_paragraph("五、材料与第三方验证计划", style="MATIC Heading 1")
    add_placeholder_table(
        doc,
        ["序号", "待补材料或验证", "提供/执行方", "用途", "优先级", "完成标准"],
        6,
    )

    doc.add_paragraph("六、工作节奏与交付", style="MATIC Heading 1")
    add_placeholder_table(
        doc,
        ["阶段", "主要工作", "触发条件", "交付物", "预计时间", "决策节点"],
        4,
    )
    doc.add_paragraph(
        "说明：日常新增材料默认更新同一详细尽调工作台；只有收到明确指令时生成阶段性或最终报告。"
    )

    props = doc.core_properties
    props.title = "医创中心详细尽调启动方案模板"
    props.subject = "医疗科技项目详细尽调启动与调查计划"
    props.author = "长三角医学先进技术创新中心"
    props.keywords = "MATIC, 详细尽调, 启动方案, 医疗科技"
    STARTUP_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    normalize_table_paragraphs(doc)
    doc.save(STARTUP_OUTPUT)
    print(STARTUP_OUTPUT)


def build_report_template():
    doc = new_matic_document(justify_body=True)
    add_cover(
        doc,
        "详细尽调报告",
        [
            ("项目名称", "[项目名称]"),
            ("报告性质", "[阶段性/最终]"),
            ("报告版本", "[V0.1/V1.0]"),
            ("当前建议", "[状态＋具体下一阶段或条件]"),
            ("证据截止日期", "[YYYY年MM月DD日]"),
            ("报告日期", "[YYYY年MM月DD日]"),
        ],
    )

    doc.add_paragraph("第一部分　详细尽调结论", style="MATIC Heading 1")
    doc.add_paragraph("一、项目简述", style="MATIC Heading 2")
    doc.add_paragraph("（一）战略性", style="MATIC Heading 3")
    doc.add_paragraph("[使用已经确认的结论进行简要概述。]")
    doc.add_paragraph("（二）先进性", style="MATIC Heading 3")
    doc.add_paragraph("[使用已经确认的结论进行简要概述。]")
    doc.add_paragraph("（三）核心优势", style="MATIC Heading 3")
    doc.add_paragraph("[概括最关键、已得到认可的核心优势。项目简述全文约300—500个中文字符，不显示评级。]")
    doc.add_page_break()

    add_strength_risk_page(doc, "二、技术方面尽调结论")
    doc.add_page_break()
    add_strength_risk_page(doc, "三、市场方面尽调结论")
    doc.add_page_break()
    add_strength_risk_page(doc, "四、核心团队尽调结论")
    doc.add_page_break()
    add_strength_risk_page(doc, "五、其他方面尽调结论", strength_count=3, risk_count=3)
    doc.add_page_break()

    doc.add_paragraph("六、AI综合研判意见", style="MATIC Heading 2")
    doc.add_paragraph(
        "[在1000个中文字符以内形成独立研判，明确当前建议、关键条件、风险提示和项目完善方向。不得只汇总BP或专家观点；新增事实须引用，推断须说明依据和边界。]"
    )
    doc.add_page_break()

    doc.add_paragraph("第二部分　详细尽调分析", style="MATIC Heading 1")
    second_part = [
        ("一、尽调范围与证据概况", "scope"),
        ("二、项目变化及问题跟踪", "issues"),
        ("三、临床需求、政策与适应证格局", "text"),
        ("四、技术专项尽调", "conclusions"),
        ("五、注册与临床转化专项", "text"),
        ("六、竞品、替代方案及相对优劣势分析", "competition"),
        ("七、市场与商业专项尽调", "market"),
        ("八、产业链与产业化专项", "text"),
        ("九、核心团队专项尽调", "conclusions"),
        ("十、知识产权、成果权属与股权专项", "conclusions"),
        ("十一、专家访谈证据与分歧", "disputes"),
        ("十二、关键风险、否决事项与前置条件", "risks"),
        ("十三、待补充材料及下一步调研计划", "plan"),
        ("十四、参考来源及附录", "sources"),
    ]
    for heading, kind in second_part:
        add_second_part_section(doc, heading, kind)

    props = doc.core_properties
    props.title = "医创中心详细尽调报告模板"
    props.subject = "医疗科技项目阶段性或最终详细尽调"
    props.author = "长三角医学先进技术创新中心"
    props.keywords = "MATIC, 详细尽调, 专家访谈, 医疗科技"
    REPORT_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    normalize_table_paragraphs(doc)
    doc.save(REPORT_OUTPUT)
    print(REPORT_OUTPUT)


def build():
    build_startup_template()
    build_report_template()


if __name__ == "__main__":
    build()
