#!/usr/bin/env python3
"""Build the formal MATIC due-diligence DOCX template."""

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "assets" / "matic-due-diligence-template.docx"

CN_BODY = "方正仿宋_GBK"
CN_TITLE = "方正小标宋_GBK"
CN_KAI = "方正楷体_GBK"
CN_HEI = "方正黑体_GBK"
EN_FONT = "Times New Roman"


def set_run_font(run, east_asia, size, bold=False):
    run.font.name = EN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), EN_FONT)
    fonts.set(qn("w:hAnsi"), EN_FONT)
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:cs"), EN_FONT)


def configure_style(style, east_asia, size, bold=False, align=None, indent=True, outline_level=None):
    style.font.name = EN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:ascii"), EN_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), EN_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(29.5)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(32) if indent else None
    if align is not None:
        pf.alignment = align
    ppr = style._element.get_or_add_pPr()
    outline = ppr.find(qn("w:outlineLvl"))
    if outline_level is not None:
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            ppr.append(outline)
        outline.set(qn("w:val"), str(outline_level))
    elif outline is not None:
        ppr.remove(outline)
    for tag in ("w:pBdr", "w:shd"):
        element = ppr.find(qn(tag))
        if element is not None:
            ppr.remove(element)


def get_or_add_paragraph_style(doc, name):
    names = [s.name for s in doc.styles]
    if name in names:
        return doc.styles[name]
    return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def add_page_number(paragraph):
    # Rebuilding a footer must replace, not append to, an existing PAGE field.
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("— ")
    set_run_font(run, CN_BODY, 10.5)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" —")
    set_run_font(run, CN_BODY, 10.5)


def enable_field_updates(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def shade_cell(cell, fill="FFFFFF"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.line_spacing = Pt(15)
    r = p.add_run(text)
    set_run_font(r, CN_BODY, 10.5, bold=bold)


def add_meta_table(doc):
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    labels = ["项目名称", "报告版本", "尽调建议", "报告日期"]
    values = ["[项目名称]", "V1.0", "初筛中", "[YYYY年MM月DD日]"]
    for row, label, value in zip(table.rows, labels, values):
        set_cell_text(row.cells[0], label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], value)
        shade_cell(row.cells[0])
        shade_cell(row.cells[1])
    return table


def add_placeholder_table(doc, headers, row_count=2):
    table = doc.add_table(rows=1 + row_count, cols=len(headers))
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in table.rows[1:]:
        for cell in row.cells:
            set_cell_text(cell, "[待填写]")
    return table


def add_project_overview_table(doc):
    labels = [
        "项目名称",
        "目标产品",
        "目标适应证、临床场景及科室",
        "核心技术",
        "当前研发阶段",
        "项目目标",
        "项目负责人",
        "实施主体或成果依托单位",
        "转化、公司化及融资状态",
        "材料与报告信息",
    ]
    table = doc.add_table(rows=len(labels), cols=2)
    table.style = "Table Grid"
    for row, label in zip(table.rows, labels):
        set_cell_text(row.cells[0], label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], "[待填写]")
    return table


def add_analysis_block(doc, heading, rating="★★★☆☆"):
    doc.add_paragraph(f"{heading}　{rating}", style="MATIC Heading 2")
    doc.add_paragraph("核心判断：[基于独立检索与核验形成一句话判断。]")
    doc.add_paragraph("主要依据：[列出关键事实并紧邻标注引用编号。]")
    doc.add_paragraph("关键不确定性：[说明尚缺证据及其对判断的影响。]")


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.2)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)

    configure_style(doc.styles["Normal"], CN_BODY, 16, indent=True)
    title_style = get_or_add_paragraph_style(doc, "MATIC Title")
    subtitle_style = get_or_add_paragraph_style(doc, "MATIC Subtitle")
    heading1_style = get_or_add_paragraph_style(doc, "MATIC Heading 1")
    heading2_style = get_or_add_paragraph_style(doc, "MATIC Heading 2")
    heading3_style = get_or_add_paragraph_style(doc, "MATIC Heading 3")
    caption_style = get_or_add_paragraph_style(doc, "MATIC Caption")
    configure_style(title_style, CN_TITLE, 22, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    configure_style(subtitle_style, CN_KAI, 16, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    configure_style(heading1_style, CN_HEI, 16, bold=True, indent=False, outline_level=0)
    configure_style(heading2_style, CN_KAI, 16, indent=False, outline_level=1)
    configure_style(heading3_style, CN_KAI, 16, indent=False, outline_level=2)
    configure_style(caption_style, CN_BODY, 10.5, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    if "Source Text" not in [s.name for s in doc.styles]:
        source = doc.styles.add_style("Source Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        source = doc.styles["Source Text"]
    configure_style(source, CN_BODY, 10.5, indent=False)

    for sec in doc.sections:
        add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(48)
    r = p.add_run("内部研判材料，仅供决策参考")
    set_run_font(r, CN_KAI, 16)

    p = doc.add_paragraph(style="MATIC Title")
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(24)
    p.add_run("[项目名称]")
    p = doc.add_paragraph(style="MATIC Subtitle")
    p.add_run("尽调报告")
    doc.add_paragraph("")
    add_meta_table(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(84)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("长三角医学先进技术创新中心")
    set_run_font(r, CN_BODY, 16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run("[YYYY年MM月DD日]")
    set_run_font(r, CN_BODY, 16)
    p.runs[0].add_break(WD_BREAK.PAGE)

    doc.add_paragraph("第一部分　项目初步研判", style="MATIC Heading 1")
    doc.add_paragraph("一、项目概要", style="MATIC Heading 2")
    doc.add_paragraph("[用约200—300字说明项目做什么、解决什么问题、拟形成什么产品及实现什么目标。BP主张须使用归因表达。]")
    add_project_overview_table(doc)
    doc.add_page_break()

    add_analysis_block(doc, "二、战略性")
    add_analysis_block(doc, "三、先进性")
    add_analysis_block(doc, "四、可行性")
    doc.add_page_break()

    doc.add_paragraph("五、AI综合研判意见", style="MATIC Heading 2")
    doc.add_paragraph("尽调建议：[状态＋拟进入的具体阶段。]")
    doc.add_paragraph("1. 核心优势：[列出2—3项经核验的项目优势。]")
    doc.add_paragraph("2. 关键风险：[列出3—5项可能影响推进的风险。]")
    doc.add_paragraph("3. 一票否决或重大前置条件：[说明确认触发、疑似触发或解除条件。]")
    doc.add_paragraph("4. 与医创中心的合作适配性：[与项目本身价值分开判断。]")
    doc.add_paragraph("5. 下一阶段及验证事项：[说明下一步动作、验证标准及停止推进条件。]")
    doc.add_page_break()

    doc.add_paragraph("第二部分　项目尽调分析", style="MATIC Heading 1")

    sections = [
        "一、项目基本情况与尽调状态",
        "二、BP关键主张核验表",
        "三、临床需求与未满足痛点",
        "四、政策环境与战略价值",
        "五、技术创新性分析",
        "六、知识产权与技术壁垒",
        "七、注册路径与临床转化可行性",
        "八、产业链分析",
        "九、国内外竞品及替代方案对比",
        "十、市场规模测算",
        "十一、团队构成分析",
        "十二、产品成熟度与产业化可行性",
        "十三、商业模式与股权架构分析",
        "十四、重大风险与一票否决事项",
        "十五、八维度五星评级及综合判断",
        "十六、与医创中心合作的参考建议",
        "十七、项目方待补充材料清单",
        "十八、进一步调研问题清单",
        "十九、参考来源及附录",
    ]
    for idx, heading in enumerate(sections):
        doc.add_paragraph(heading, style="MATIC Heading 2")
        if idx == 1:
            add_placeholder_table(doc, ["序号", "BP关键主张", "外部证据", "核验结果", "置信度", "影响"], 3)
        elif idx == 8:
            add_placeholder_table(doc, ["方案", "类型", "注册/应用状态", "关键差异", "证据来源"], 3)
        elif idx == 13:
            add_placeholder_table(doc, ["事项", "触发等级", "事实与证据", "影响及解除条件"], 2)
        elif idx == 16:
            add_placeholder_table(doc, ["序号", "待补充材料", "用途", "优先级"], 3)
        elif idx == 17:
            add_placeholder_table(doc, ["对象", "问题", "拟验证事项"], 3)
        elif idx == 18:
            add_placeholder_table(doc, ["序号", "来源", "标题", "日期", "链接"], 3)
        else:
            doc.add_paragraph("[根据BP、公开检索和独立分析填写。区分项目方陈述、公开事实、推断与待核实事项。]")

    props = doc.core_properties
    props.title = "医创中心项目尽调报告模板"
    props.subject = "医疗科技项目BP核验与内部尽调"
    props.author = "长三角医学先进技术创新中心"
    props.keywords = "MATIC, 尽调, 医疗科技, BP"
    enable_field_updates(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
